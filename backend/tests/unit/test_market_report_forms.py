"""Report forms of the Vietnamese and American markets.

The Reports page offered BOQ, cost and GAEB exports and nothing a Vietnamese
or American project hands to an owner. These workbooks arrange figures the
platform already computes into the market's own forms: the TT 11/2021
estimate summary and Phụ lục 03a for Vietnam, AIA G702/G703 for the US.
"""

from __future__ import annotations

import io
from decimal import Decimal

from openpyxl import load_workbook

from app.modules.reporting.market_forms import (
    ClaimContext,
    EstimateSummary,
    MarkupLine,
    build_estimate_summary_rows,
    forms_for_country,
    market_for_country,
    render_aia_xlsx,
    render_estimate_summary_xlsx,
    render_payment_03a_xlsx,
)

D = Decimal


def _summary() -> EstimateSummary:
    return EstimateSummary(
        project_name="Nhà xưởng Bình Dương",
        boq_name="Dự toán phần thô",
        currency="VND",
        direct_cost=D("1000000000"),
        by_type={"material": D("600000000"), "labor": D("250000000"), "equipment": D("100000000")},
        markups=[
            MarkupLine("Chi phí chung", "overhead", D("6.5"), D("65000000")),
            MarkupLine("Chi phí nhà tạm", "overhead", D("1.1"), D("11000000")),
            MarkupLine("Thu nhập chịu thuế tính trước", "profit", D("6"), D("64560000"), "cumulative"),
            MarkupLine("Thuế GTGT", "tax", D("10"), D("114056000.4"), "cumulative"),
        ],
    )


def _sheet_values(content: bytes, sheet: str | None = None) -> list[list[object]]:
    wb = load_workbook(io.BytesIO(content))
    ws = wb[sheet] if sheet else wb.active
    return [list(row) for row in ws.iter_rows(values_only=True)]


def test_each_market_gets_its_own_forms() -> None:
    assert market_for_country("vn") == "VN"
    assert market_for_country("US") == "US"
    assert [f.id for f in forms_for_country("VN")] == ["vn_estimate_summary", "vn_payment_03a"]
    assert [f.id for f in forms_for_country("US")] == ["aia_g702_g703"]
    assert len(forms_for_country(None)) == 3


def test_the_estimate_summary_follows_the_circular_and_adds_up() -> None:
    rows = {r.symbol: r.value for r in build_estimate_summary_rows(_summary()) if r.symbol}
    assert rows["VL"] == D("600000000")
    assert rows["NC"] == D("250000000")
    assert rows["M"] == D("100000000")
    # What the resource breakdown does not classify stays inside T.
    assert rows["K"] == D("50000000")
    assert rows["T"] == D("1000000000")
    assert rows["GT"] == D("76000000")
    assert rows["TL"] == D("64560000")
    assert rows["G"] == D("1140560000")
    # VND rounds to whole dong, once.
    assert rows["GTGT"] == D("114056000")
    assert rows["GXD"] == D("1254616000")


def test_the_estimate_summary_workbook_names_the_circular_and_prints_whole_dong() -> None:
    values = _sheet_values(render_estimate_summary_xlsx(_summary()))
    text = " ".join(str(v) for row in values for v in row if v is not None)
    assert "BẢNG TỔNG HỢP DỰ TOÁN CHI PHÍ XÂY DỰNG" in text
    assert "11/2021/TT-BXD" in text
    gxd = next(row for row in values if "GXD" in row)
    assert gxd[3] == 1254616000


def _application() -> dict[str, object]:
    return {
        "contract_id": "c1",
        "application_number": "03",
        "period_start": None,
        "period_end": None,
        "claim_date": None,
        "currency": "USD",
        "summary": {
            "original_contract_sum": D("100000"),
            "change_orders_net": D("5000"),
            "contract_sum_to_date": D("105000"),
            "total_completed_stored": D("42000"),
            "retainage": D("4200"),
            "total_earned_less_retainage": D("37800"),
            "previous_certificates_total": D("18000"),
            "current_payment_due": D("19800"),
            "balance_to_finish": D("67200"),
        },
        "lines": [
            {
                "item_number": "1",
                "description": "Foundations",
                "scheduled_value": D("60000"),
                "previous_value": D("20000"),
                "this_period_value": D("10000"),
                "materials_stored": D("0"),
                "total_completed_stored": D("30000"),
                "percent_complete": D("50.00"),
                "balance_to_finish": D("30000"),
                "retainage": D("3000"),
            },
            {
                "item_number": "2",
                "description": "Frame",
                "scheduled_value": D("45000"),
                "previous_value": D("0"),
                "this_period_value": D("12000"),
                "materials_stored": D("0"),
                "total_completed_stored": D("12000"),
                "percent_complete": D("26.67"),
                "balance_to_finish": D("33000"),
                "retainage": D("1200"),
            },
        ],
    }


def test_appendix_03a_carries_the_claim_figures_and_the_amount_due() -> None:
    payload = _application()
    payload["currency"] = "VND"
    ctx = ClaimContext(project_name="Dự án A", contract_title="Thi công phần thô", contract_code="HĐ-01")
    values = _sheet_values(render_payment_03a_xlsx(payload, ctx))
    text = " ".join(str(v) for row in values for v in row if v is not None)
    assert "Phụ lục 03a" in text
    total = next(row for row in values if row[1] == "Tổng cộng")
    assert total[2:6] == [105000, 20000, 22000, 42000]
    due = next(row for row in values if row[1] == "Giá trị đề nghị thanh toán lần này")
    assert due[5] == 19800


def test_the_aia_workbook_has_the_summary_and_the_continuation_sheet() -> None:
    ctx = ClaimContext(project_name="Austin Office", contract_title="Shell and core", contract_code="C-7")
    content = render_aia_xlsx(_application(), ctx)
    g702 = _sheet_values(content, "G702")
    line8 = next(row for row in g702 if row[0] == "8")
    assert line8[2] == 19800
    g703 = _sheet_values(content, "G703")
    foundations = next(row for row in g703 if row[1] == "Foundations")
    assert foundations[2] == 60000 and foundations[6] == 30000 and foundations[9] == 3000
