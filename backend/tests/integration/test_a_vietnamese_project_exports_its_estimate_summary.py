# DDC-CWICR-OE: DataDrivenConstruction · OpenConstructionERP
# Copyright (c) 2026 Artem Boiko / DataDrivenConstruction
"""A Vietnamese project's Reports page offers its market's forms, and they download.

End to end through the reporting routes: the form list follows the project's
country, and the TT 11/2021 estimate summary of a real bill opens as a
workbook whose direct cost is the bill's.
"""

from __future__ import annotations

import asyncio
import io
import uuid

import pytest
import pytest_asyncio
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook

from app.main import create_app


@pytest_asyncio.fixture(scope="module")
async def client():
    app = create_app()
    async with app.router.lifespan_context(app):
        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://test") as ac:
            yield ac


@pytest_asyncio.fixture(scope="module")
async def auth(client: AsyncClient) -> dict[str, str]:
    unique = uuid.uuid4().hex[:8]
    email = f"vnforms-{unique}@test.io"
    password = f"VnForms{unique}9"
    reg = await client.post(
        "/api/v1/users/auth/register",
        json={"email": email, "password": password, "full_name": "VN Forms Tester", "role": "admin"},
    )
    assert reg.status_code == 201, reg.text
    from ._auth_helpers import promote_to_admin

    await promote_to_admin(email)
    token = ""
    for attempt in range(3):
        resp = await client.post("/api/v1/users/auth/login", json={"email": email, "password": password})
        token = resp.json().get("access_token", "")
        if token:
            break
        await asyncio.sleep(5 * (attempt + 1))
    assert token
    return {"Authorization": f"Bearer {token}"}


@pytest.mark.asyncio
async def test_the_estimate_summary_downloads_for_a_vietnamese_bill(client: AsyncClient, auth: dict[str, str]) -> None:
    resp = await client.post(
        "/api/v1/projects/",
        json={
            "name": f"Nhà xưởng {uuid.uuid4().hex[:6]}",
            "description": "",
            "region": "VN",
            "country_code": "VN",
            "classification_standard": "masterformat",
            "currency": "VND",
            "locale": "vi",
        },
        headers=auth,
    )
    assert resp.status_code == 201, resp.text
    project = resp.json()

    resp = await client.get(f"/api/v1/reporting/market-forms/?project_id={project['id']}", headers=auth)
    assert resp.status_code == 200, resp.text
    assert resp.json()["market"] == "VN"
    assert [f["id"] for f in resp.json()["forms"]] == ["vn_estimate_summary", "vn_payment_03a"]
    assert resp.json()["claims"] == []

    resp = await client.post(
        "/api/v1/boq/boqs/", json={"project_id": project["id"], "name": "Phần thô", "description": ""}, headers=auth
    )
    assert resp.status_code == 201, resp.text
    boq_id = resp.json()["id"]
    resp = await client.post(
        f"/api/v1/boq/boqs/{boq_id}/positions/",
        json={
            "boq_id": boq_id,
            "ordinal": "01",
            "description": "Bê tông móng",
            "unit": "m3",
            "quantity": 10,
            "unit_rate": 1250000,
        },
        headers=auth,
    )
    assert resp.status_code == 201, resp.text

    resp = await client.get(f"/api/v1/reporting/market-forms/vn-estimate-summary/?boq_id={boq_id}", headers=auth)
    assert resp.status_code == 200, resp.text[:300]
    rows = [list(r) for r in load_workbook(io.BytesIO(resp.content)).active.iter_rows(values_only=True)]
    direct = next(r for r in rows if "T" in r and r[1] == "Cộng chi phí trực tiếp")
    assert direct[3] == 12500000


@pytest.mark.asyncio
async def test_the_forms_need_a_signed_in_caller(client: AsyncClient) -> None:
    resp = await client.get(f"/api/v1/reporting/market-forms/vn-estimate-summary/?boq_id={uuid.uuid4()}")
    assert resp.status_code == 401
