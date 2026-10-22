# DDC-CWICR-OE: DataDrivenConstruction · OpenConstructionERP
# Copyright (c) 2026 Artem Boiko / DataDrivenConstruction
"""Statutory and customary report forms of the markets this build serves.

Three workbooks, each drawn from figures the platform already computes, so no
number is worked out a second time here:

* ``vn_estimate_summary``: Vietnam, Bảng tổng hợp dự toán chi phí xây dựng in
  the shape Thông tư 11/2021/TT-BXD (Phụ lục III) gives it. Direct cost split
  into materials (VL), labour (NC) and plant (M) from the bill's resource
  breakdown, then the bill's own markup lines in the circular's order:
  indirect cost (GT), pre-determined taxable income (TL), VAT. The rates are
  the bill's; the form only arranges them.
* ``vn_payment_03a``: Vietnam, Phụ lục 03a, Bảng xác định giá trị khối lượng
  công việc hoàn thành theo hợp đồng đề nghị thanh toán, from a progress claim
  (the same country-independent figures the AIA continuation sheet reads).
* ``aia_g702_g703``: United States, AIA G702 summary and G703 continuation
  sheet as a workbook, from the contracts module's AIA builder. The PDF of the
  same application is served by the contracts module.

Amounts are Decimal end to end and rounded once, to the currency's minor
units, when a cell is written (VND to whole dong, USD to cents).
"""

from __future__ import annotations

import io
from collections.abc import Iterable, Sequence
from dataclasses import dataclass, replace
from datetime import date
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from typing import Any

from app.core.money import minor_units, money_quantum

ZERO = Decimal("0")

#: Country -> market whose forms apply. The AIA forms follow the contracts
#: module's eligibility (US, CA, AU).
_MARKET_BY_COUNTRY = {"VN": "VN", "US": "US", "CA": "US", "AU": "US"}


@dataclass(frozen=True, slots=True)
class MarketForm:
    """One form the Reports page can offer."""

    id: str
    market: str
    #: What the form is generated from: a bill ("boq") or a progress claim ("claim").
    source: str
    formats: tuple[str, ...]


FORMS: tuple[MarketForm, ...] = (
    MarketForm("vn_estimate_summary", "VN", "boq", ("xlsx",)),
    MarketForm("vn_payment_03a", "VN", "claim", ("xlsx",)),
    MarketForm("aia_g702_g703", "US", "claim", ("xlsx", "pdf")),
)


def market_for_country(country: str | None) -> str | None:
    """The market whose forms a project in ``country`` uses, if this build has any."""
    return _MARKET_BY_COUNTRY.get((country or "").strip().upper())


def forms_for_country(country: str | None) -> list[MarketForm]:
    """Forms of the project's market; every form when the market is unknown."""
    market = market_for_country(country)
    if market is None:
        return list(FORMS)
    return [f for f in FORMS if f.market == market]


def _dec(value: Any) -> Decimal:
    if value is None or value == "":
        return ZERO
    try:
        return Decimal(str(value))
    except (InvalidOperation, ValueError):
        return ZERO


def _round(value: Decimal, currency: str) -> Decimal:
    return value.quantize(money_quantum(currency), rounding=ROUND_HALF_UP)


# ── Bảng tổng hợp dự toán (TT 11/2021) ──────────────────────────────────


@dataclass(frozen=True, slots=True)
class MarkupLine:
    """One markup line of a bill, as the BOQ engine computed it."""

    name: str
    category: str
    percentage: Decimal
    amount: Decimal
    apply_to: str = "direct_cost"


@dataclass(frozen=True, slots=True)
class SummaryRow:
    """One row of the estimate summary."""

    number: str
    content: str
    basis: str
    value: Decimal | None
    symbol: str
    bold: bool = False


@dataclass(slots=True)
class EstimateSummary:
    """The figures a Bảng tổng hợp dự toán is drawn from."""

    project_name: str
    boq_name: str
    currency: str
    direct_cost: Decimal
    #: Direct cost by resource type as the BOQ cost breakdown names them.
    by_type: dict[str, Decimal]
    markups: list[MarkupLine]
    tax_rate: Decimal | None = None
    tax_amount: Decimal = ZERO


def _pct(value: Decimal) -> str:
    """A rate as Vietnamese forms print it: no trailing zeros, decimal comma."""
    return format(value.normalize(), "f").replace(".", ",")


def build_estimate_summary_rows(summary: EstimateSummary) -> list[SummaryRow]:
    """Arrange the bill's figures in the circular's order.

    VL is material, NC labour, M machinery and equipment. Subcontracted and
    unclassified cost stays inside T on its own row, so T is the bill's direct
    cost to the dong. Markups are placed by their category: ``tax`` under VAT,
    ``profit`` as TL, everything else (overhead, contingency, site
    facilities) under indirect cost.
    """
    cur = summary.currency
    vl = summary.by_type.get("material", ZERO)
    nc = summary.by_type.get("labor", ZERO)
    m = summary.by_type.get("machinery", ZERO) + summary.by_type.get("equipment", ZERO)
    other = summary.direct_cost - vl - nc - m
    t = summary.direct_cost

    rows: list[SummaryRow] = [
        SummaryRow("I", "CHI PHÍ TRỰC TIẾP", "", None, "", bold=True),
        SummaryRow("1", "Chi phí vật liệu", "Theo bảng dự toán", vl, "VL"),
        SummaryRow("2", "Chi phí nhân công", "Theo bảng dự toán", nc, "NC"),
        SummaryRow("3", "Chi phí máy và thiết bị thi công", "Theo bảng dự toán", m, "M"),
    ]
    if other != ZERO:
        rows.append(SummaryRow("4", "Chi phí khác (thầu phụ, chưa phân loại)", "Theo bảng dự toán", other, "K"))
    rows.append(
        SummaryRow("", "Cộng chi phí trực tiếp", "VL + NC + M" + (" + K" if other != ZERO else ""), t, "T", True)
    )

    indirect = [mk for mk in summary.markups if mk.category not in ("profit", "tax")]
    profit = [mk for mk in summary.markups if mk.category == "profit"]
    taxes = [mk for mk in summary.markups if mk.category == "tax"]

    def basis(mk: MarkupLine, base_symbol: str) -> str:
        if mk.percentage == ZERO:
            return "Giá trị cố định"
        return f"{base_symbol} x {_pct(mk.percentage)}%"

    gt = sum((mk.amount for mk in indirect), ZERO)
    rows.append(SummaryRow("II", "CHI PHÍ GIÁN TIẾP", "", None, "", bold=True))
    for i, mk in enumerate(indirect, start=1):
        rows.append(
            SummaryRow(str(i), mk.name, basis(mk, "T" if mk.apply_to == "direct_cost" else "T + GT"), mk.amount, "")
        )
    rows.append(SummaryRow("", "Cộng chi phí gián tiếp", "", gt, "GT", True))

    tl = sum((mk.amount for mk in profit), ZERO)
    rows.append(SummaryRow("III", "THU NHẬP CHỊU THUẾ TÍNH TRƯỚC", "", None, "", bold=True))
    for i, mk in enumerate(profit, start=1):
        rows.append(SummaryRow(str(i), mk.name, basis(mk, "(T + GT)"), mk.amount, ""))
    rows.append(SummaryRow("", "Cộng thu nhập chịu thuế tính trước", "", tl, "TL", True))

    g = t + gt + tl
    rows.append(SummaryRow("", "Chi phí xây dựng trước thuế", "T + GT + TL", g, "G", True))

    vat = sum((mk.amount for mk in taxes), ZERO) + summary.tax_amount
    rows.append(SummaryRow("IV", "THUẾ GIÁ TRỊ GIA TĂNG", "", None, "", bold=True))
    for i, mk in enumerate(taxes, start=1):
        rows.append(SummaryRow(str(i), mk.name, basis(mk, "G"), mk.amount, ""))
    if summary.tax_amount != ZERO:
        rate = f"G x {_pct(summary.tax_rate)}%" if summary.tax_rate else ""
        rows.append(SummaryRow(str(len(taxes) + 1), "Thuế GTGT", rate, summary.tax_amount, ""))
    rows.append(SummaryRow("", "Cộng thuế giá trị gia tăng", "", vat, "GTGT", True))
    rows.append(SummaryRow("", "CHI PHÍ XÂY DỰNG SAU THUẾ", "G + GTGT", g + vat, "GXD", True))
    if not cur:
        return rows
    return [row if row.value is None else replace(row, value=_round(row.value, cur)) for row in rows]


# ── Workbook helpers ────────────────────────────────────────────────────


def _number_format(currency: str) -> str:
    digits = minor_units(currency) if currency else 2
    return "#,##0" if digits == 0 else "#,##0." + "0" * digits


def _workbook() -> tuple[Any, Any]:
    from openpyxl import Workbook

    wb = Workbook()
    return wb, wb.active


def _finish(wb: Any, ws: Any, *, title: str, details: Sequence[str]) -> bytes:
    from app.core.xlsx_branding import apply_company_header
    from app.core.xlsx_text import store_strings_as_text

    store_strings_as_text(ws)
    apply_company_header(ws, title=title, details=list(details))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _header(ws: Any, row: int, labels: Iterable[str]) -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    thin = Side(style="thin", color="999999")
    for col, label in enumerate(labels, start=1):
        cell = ws.cell(row=row, column=col, value=label)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.fill = PatternFill("solid", fgColor="EEF2F7")
        cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)


def _widths(ws: Any, widths: Sequence[float]) -> None:
    from openpyxl.utils import get_column_letter

    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width


def render_estimate_summary_xlsx(summary: EstimateSummary) -> bytes:
    """The Bảng tổng hợp dự toán chi phí xây dựng workbook."""
    from openpyxl.styles import Font

    rows = build_estimate_summary_rows(summary)
    wb, ws = _workbook()
    ws.title = "Tong hop du toan"
    _widths(ws, (8, 52, 22, 20, 10))
    _header(ws, 1, ("STT", "Nội dung chi phí", "Cách tính", f"Giá trị ({summary.currency})", "Ký hiệu"))
    fmt = _number_format(summary.currency)
    for r, row in enumerate(rows, start=2):
        values = (row.number, row.content, row.basis, row.value, row.symbol)
        for c, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            if row.bold:
                cell.font = Font(bold=True)
            if c == 4 and value is not None:
                cell.number_format = fmt
    ws.freeze_panes = "A2"
    return _finish(
        wb,
        ws,
        title="BẢNG TỔNG HỢP DỰ TOÁN CHI PHÍ XÂY DỰNG",
        details=(
            "Theo Thông tư 11/2021/TT-BXD (Phụ lục III)",
            f"Dự án: {summary.project_name}",
            f"Dự toán: {summary.boq_name}",
            f"Đơn vị tiền tệ: {summary.currency}",
        ),
    )


# ── Phụ lục 03a and AIA G702/G703, from a payment application ───────────


@dataclass(slots=True)
class ClaimContext:
    """Names around a payment application."""

    project_name: str
    contract_title: str
    contract_code: str
    counterparty: str = ""


def _period(payload: dict[str, Any]) -> str:
    start, end = payload.get("period_start"), payload.get("period_end")

    def text(value: Any) -> str:
        if isinstance(value, date):
            return value.strftime("%d/%m/%Y")
        return str(value or "")

    if start or end:
        return f"{text(start)} - {text(end)}"
    return text(payload.get("claim_date"))


def render_payment_03a_xlsx(payload: dict[str, Any], ctx: ClaimContext) -> bytes:
    """Phụ lục 03a from the figures of one progress claim.

    Value columns come from the claim's schedule-of-values lines: the contract
    value, the value done up to the previous payment, in this period and in
    total. Quantities are not tracked on claim lines, so the volume columns
    carry the share of each line completed.
    """
    from openpyxl.styles import Font

    currency = str(payload.get("currency") or "")
    lines: list[dict[str, Any]] = list(payload.get("lines") or [])
    summary: dict[str, Any] = dict(payload.get("summary") or {})
    fmt = _number_format(currency)

    wb, ws = _workbook()
    ws.title = "Phu luc 03a"
    _widths(ws, (6, 46, 18, 18, 18, 18, 12, 18))
    _header(
        ws,
        1,
        (
            "TT",
            "Tên công việc",
            f"Giá trị theo hợp đồng ({currency})",
            "Lũy kế đến hết lần thanh toán trước",
            "Thực hiện lần này",
            "Lũy kế đến hết lần này",
            "Tỷ lệ hoàn thành (%)",
            "Còn lại theo hợp đồng",
        ),
    )
    r = 2
    totals = {"scheduled": ZERO, "previous": ZERO, "this": ZERO, "total": ZERO}
    for i, line in enumerate(lines, start=1):
        scheduled = line.get("scheduled_value")
        previous = _dec(line.get("previous_value"))
        this_period = _dec(line.get("this_period_value")) + _dec(line.get("materials_stored"))
        total = _dec(line.get("total_completed_stored"))
        totals["scheduled"] += _dec(scheduled)
        totals["previous"] += previous
        totals["this"] += this_period
        totals["total"] += total
        values: list[Any] = [
            str(line.get("item_number") or i),
            str(line.get("description") or ""),
            _round(_dec(scheduled), currency) if scheduled is not None else None,
            _round(previous, currency),
            _round(this_period, currency),
            _round(total, currency),
            _dec(line.get("percent_complete")) if line.get("percent_complete") is not None else None,
            _round(_dec(line.get("balance_to_finish")), currency)
            if line.get("balance_to_finish") is not None
            else None,
        ]
        for c, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            if c in (3, 4, 5, 6, 8) and value is not None:
                cell.number_format = fmt
            if c == 7 and value is not None:
                cell.number_format = "0.00"
        r += 1

    total_row = ["", "Tổng cộng", *(_round(totals[k], currency) for k in ("scheduled", "previous", "this", "total"))]
    for c, value in enumerate(total_row, start=1):
        cell = ws.cell(row=r, column=c, value=value)
        cell.font = Font(bold=True)
        if c >= 3:
            cell.number_format = fmt
    r += 2

    footer = (
        ("Giá trị hợp đồng (kể cả điều chỉnh)", summary.get("contract_sum_to_date")),
        ("Lũy kế giá trị khối lượng thực hiện đến hết lần này", summary.get("total_completed_stored")),
        ("Giá trị giữ lại (bảo lưu) lũy kế", summary.get("retainage")),
        ("Lũy kế giá trị đề nghị thanh toán sau khi giữ lại", summary.get("total_earned_less_retainage")),
        ("Số tiền đã thanh toán đến hết lần trước", summary.get("previous_certificates_total")),
        ("Giá trị đề nghị thanh toán lần này", summary.get("current_payment_due")),
    )
    for label, value in footer:
        ws.cell(row=r, column=2, value=label).font = Font(bold=label.startswith("Giá trị đề nghị"))
        cell = ws.cell(row=r, column=6, value=_round(_dec(value), currency))
        cell.number_format = fmt
        r += 1
    r += 1
    ws.cell(row=r, column=2, value="ĐẠI DIỆN BÊN GIAO THẦU").font = Font(bold=True)
    ws.cell(row=r, column=6, value="ĐẠI DIỆN BÊN NHẬN THẦU").font = Font(bold=True)

    return _finish(
        wb,
        ws,
        title="BẢNG XÁC ĐỊNH GIÁ TRỊ KHỐI LƯỢNG CÔNG VIỆC HOÀN THÀNH THEO HỢP ĐỒNG ĐỀ NGHỊ THANH TOÁN",
        details=(
            "Phụ lục 03a",
            f"Dự án: {ctx.project_name}",
            f"Hợp đồng: {ctx.contract_code} {ctx.contract_title}".strip(),
            *([f"Bên nhận thầu: {ctx.counterparty}"] if ctx.counterparty else []),
            f"Lần thanh toán: {payload.get('application_number') or ''}",
            f"Giai đoạn: {_period(payload)}",
            f"Đơn vị tiền tệ: {currency}",
        ),
    )


def render_aia_xlsx(payload: dict[str, Any], ctx: ClaimContext) -> bytes:
    """AIA G702 (summary) and G703 (continuation sheet) as one workbook."""
    from openpyxl.styles import Font

    currency = str(payload.get("currency") or "")
    fmt = _number_format(currency)
    summary: dict[str, Any] = dict(payload.get("summary") or {})
    wb, g702 = _workbook()
    g702.title = "G702"
    _widths(g702, (8, 58, 22))
    _header(g702, 1, ("Line", "Application and Certificate for Payment", f"Amount ({currency})"))
    g702_rows = (
        ("1", "Original contract sum", "original_contract_sum"),
        ("2", "Net change by change orders", "change_orders_net"),
        ("3", "Contract sum to date (line 1 + 2)", "contract_sum_to_date"),
        ("4", "Total completed and stored to date (G703 column G)", "total_completed_stored"),
        ("5", "Retainage", "retainage"),
        ("6", "Total earned less retainage (line 4 - line 5)", "total_earned_less_retainage"),
        ("7", "Less previous certificates for payment", "previous_certificates_total"),
        ("8", "Current payment due", "current_payment_due"),
        ("9", "Balance to finish, including retainage", "balance_to_finish"),
    )
    for r, (line, label, key) in enumerate(g702_rows, start=2):
        g702.cell(row=r, column=1, value=line)
        g702.cell(row=r, column=2, value=label).font = Font(bold=key == "current_payment_due")
        cell = g702.cell(row=r, column=3, value=_round(_dec(summary.get(key)), currency))
        cell.number_format = fmt

    g703 = wb.create_sheet("G703")
    _widths(g703, (8, 40, 16, 16, 16, 16, 16, 10, 16, 16))
    _header(
        g703,
        1,
        (
            "A Item",
            "B Description of work",
            "C Scheduled value",
            "D From previous application",
            "E This period",
            "F Materials presently stored",
            "G Total completed and stored",
            "% (G / C)",
            "H Balance to finish",
            "I Retainage",
        ),
    )
    r = 2
    for line in payload.get("lines") or []:
        values: list[Any] = [
            str(line.get("item_number") or ""),
            str(line.get("description") or ""),
            *(
                _round(_dec(line.get(k)), currency) if line.get(k) is not None else None
                for k in (
                    "scheduled_value",
                    "previous_value",
                    "this_period_value",
                    "materials_stored",
                    "total_completed_stored",
                )
            ),
            _dec(line.get("percent_complete")) if line.get("percent_complete") is not None else None,
            _round(_dec(line.get("balance_to_finish")), currency)
            if line.get("balance_to_finish") is not None
            else None,
            _round(_dec(line.get("retainage")), currency),
        ]
        for c, value in enumerate(values, start=1):
            cell = g703.cell(row=r, column=c, value=value)
            if value is not None and c in (3, 4, 5, 6, 7, 9, 10):
                cell.number_format = fmt
            if value is not None and c == 8:
                cell.number_format = "0.00"
        r += 1

    from app.core.xlsx_branding import apply_company_header
    from app.core.xlsx_text import store_strings_as_text

    store_strings_as_text(g703)
    apply_company_header(g703, title="CONTINUATION SHEET (AIA G703)", details=[f"Project: {ctx.project_name}"])
    return _finish(
        wb,
        g702,
        title="APPLICATION AND CERTIFICATE FOR PAYMENT (AIA G702)",
        details=(
            f"Project: {ctx.project_name}",
            f"Contract: {ctx.contract_code} {ctx.contract_title}".strip(),
            f"Application no.: {payload.get('application_number') or ''}",
            f"Period: {_period(payload)}",
            f"Currency: {currency}",
        ),
    )


__all__ = [
    "FORMS",
    "ClaimContext",
    "EstimateSummary",
    "MarketForm",
    "MarkupLine",
    "SummaryRow",
    "build_estimate_summary_rows",
    "forms_for_country",
    "market_for_country",
    "render_aia_xlsx",
    "render_estimate_summary_xlsx",
    "render_payment_03a_xlsx",
]
